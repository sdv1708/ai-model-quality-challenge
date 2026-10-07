"""Independent browser-test data; does not implement the learner's backend factory."""

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

HEADERS = [
    "Input Length",
    "Output Length",
    "Cache %",
    "Batch Size",
    "Max number of milliseconds",
    "Target Max number of milliseconds",
    "Prompt only Throughput (t/s)",
    "Gen only Throughput (t/s)",
    "Throughput (t/s)",
    "Throughput / box (t/s/hardware)",
    "Uncached Throughput (t/s)",
    "Uncached Throughput / box (t/s/hardware)",
    "Cached Throughput (t/s)",
    "Cached Throughput / box (t/s/hardware)",
    "TTFT (ms)",
    "Real Prompt Speed (t/s/user)",
    "Prompt Speed with Queueing (t/s/user)",
    "Gen Speed (t/s/user)",
    "RPM",
]


def make_workbook(multiplier: int) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Summary"
    sheet.append(["Synthetic browser-test projections"])
    sheet.append(HEADERS)
    for batch, prompt, generation, ttft, speed in [
        (10, 150, 110, 12.5, 1350),
        (20, 220, 180, 25, 1220),
    ]:
        throughput = (prompt + generation) * multiplier
        sheet.append(
            [
                100,
                100,
                0.5,
                batch,
                500,
                600,
                prompt * multiplier,
                generation * multiplier,
                throughput,
                throughput / 10,
                throughput / 2,
                throughput / 20,
                throughput / 2,
                throughput / 20,
                ttft / multiplier,
                1500 * multiplier,
                1400 * multiplier,
                speed * multiplier,
                throughput * 60 / 200,
            ]
        )
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def mutate(original: bytes, case: str) -> bytes:
    workbook = load_workbook(BytesIO(original))
    sheet = workbook["Summary"]
    if case == "missing-column":
        sheet["D2"] = None
    elif case == "invalid-number":
        sheet["D3"] = "not-a-number"
    elif case == "empty-table":
        sheet.delete_rows(3, 2)
    else:
        raise ValueError(case)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def main() -> None:
    model_l = make_workbook(2)
    files = {
        "Model A profile 1.xlsx": make_workbook(1),
        "Model L profile 1.xlsx": model_l,
        **{
            f"{case}.xlsx": mutate(model_l, case)
            for case in ["missing-column", "invalid-number", "empty-table"]
        },
        "corrupt.xlsx": b"Not an Excel archive.",
    }
    destination = Path(__file__).resolve().parents[1] / "fixtures" / "browser-generated"
    destination.mkdir(parents=True, exist_ok=True)
    for filename, contents in files.items():
        (destination / filename).write_bytes(contents)


if __name__ == "__main__":
    main()

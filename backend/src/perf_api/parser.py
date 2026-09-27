import io
import re
from typing import Any

import openpyxl  # type: ignore[import-untyped]

from perf_api.schemas import PerformanceRecord, WorkbookNormalizationResponse

EXPECTED_COLUMNS = [
    "Input Length",
    "Output Length",
    "Cache %",
    "Batch Size",
    "Max number of milliseconds",
    "Target Max number of milliseconds",
    "Prompt only Throughput (t/s)",
    "Gen only Throughput (t/s)",
    "Throughput (t/s)",
    "Throughput / box (t/s/hardware)",
    "Uncached Throughput (t/s)",
    "Uncached Throughput / box (t/s/hardware)",
    "Cached Throughput (t/s)",
    "Cached Throughput / box (t/s/hardware)",
    "TTFT (ms)",
    "Real Prompt Speed (t/s/user)",
    "Prompt Speed with Queueing (t/s/user)",
    "Gen Speed (t/s/user)",
    "RPM",
]


class WorkbookParseError(ValueError):
    """Raised when a workbook fails structure or numeric validation."""


def parse_and_normalize_workbook(file_bytes: bytes, filename: str) -> WorkbookNormalizationResponse:
    """Parse openpyxl workbook bytes and normalize merged cells into typed records."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as err:
        raise WorkbookParseError(f"Invalid Excel workbook file: {err}") from err

    if "Summary" not in wb.sheetnames:
        raise WorkbookParseError("Workbook missing required 'Summary' sheet.")

    sheet = wb["Summary"]

    # Extract headers from Row 2
    header_row_idx = 2
    headers = [sheet.cell(row=header_row_idx, column=c).value for c in range(1, 20)]
    missing_cols = [col for col in EXPECTED_COLUMNS if col not in (headers or [])]
    if missing_cols:
        raise WorkbookParseError(f"Missing required columns in 'Summary' sheet: {missing_cols}")

    # Build column index mapping
    col_map = {col: headers.index(col) + 1 for col in EXPECTED_COLUMNS}

    records: list[PerformanceRecord] = []
    current_forward_fills: dict[str, Any] = {
        "Input Length": None,
        "Output Length": None,
        "Cache %": None,
    }

    # Iterate data rows starting at Row 3
    row_idx = 3
    while True:
        batch_size_val = sheet.cell(row=row_idx, column=col_map["Batch Size"]).value
        if batch_size_val is None:
            if any(
                sheet.cell(row=row_idx, column=col_map[col]).value is not None
                for col in EXPECTED_COLUMNS
            ):
                raise WorkbookParseError(
                    f"Null value encountered at row {row_idx}, column 'Batch Size'"
                )
            if any(
                sheet.cell(row=later_row, column=column).value is not None
                for later_row in range(row_idx + 1, sheet.max_row + 1)
                for column in col_map.values()
            ):
                raise WorkbookParseError(f"Unexpected blank row at row {row_idx} before more data")
            break  # Reached end of data table

        row_data: dict[str, Any] = {}
        for col_name in EXPECTED_COLUMNS:
            val = sheet.cell(row=row_idx, column=col_map[col_name]).value

            # Forward-fill merged columns (Input Length, Output Length, Cache %)
            if col_name in current_forward_fills:
                if val is not None:
                    current_forward_fills[col_name] = val
                else:
                    val = current_forward_fills[col_name]

            if val is None:
                raise WorkbookParseError(
                    f"Null value encountered at row {row_idx}, column {col_name!r}"
                )
            row_data[col_name] = val

        try:
            records.append(PerformanceRecord.model_validate(row_data))
        except Exception as err:
            raise WorkbookParseError(f"Row {row_idx} numeric validation failed: {err}") from err

        row_idx += 1

    if not records:
        raise WorkbookParseError("Workbook contained no performance data rows.")

    # Extract Model Name and Profile ID dynamically from filename
    # e.g., "Model A profile 1.xlsx" -> model_name="Model A", profile_id="1"
    match = re.search(r"(Model\s+[A-Z0-9_-]+)\s+profile\s+(\d+)", filename, re.IGNORECASE)
    if match:
        model_name, profile_id = match.group(1), match.group(2)
    else:
        model_name, profile_id = filename.split(".")[0], "unknown"

    return WorkbookNormalizationResponse(
        model_name=model_name,
        profile_id=profile_id,
        record_count=len(records),
        records=records,
    )

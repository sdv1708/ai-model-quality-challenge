"""Fresh issue #10 workbook fixtures and optional export command.

Run from backend: uv run python tests/resilience_fixtures.py
These are learner-owned API fixtures. Frontend tests generate independent inputs.
"""

import io
from pathlib import Path
from typing import Literal

from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]

HEADERS = [
    "Input Length",
    "Output Length",
    "Cache %",
    "Batch Size",
    "Max number of milliseconds",
    "Target Max number of milliseconds",
    "Prompt only Throughput (t/s)",
    "Gen only Throughput (t/s)",
    "Throughput (t/s)",
    "Throughput / box (t/s/hardware)",
    "Uncached Throughput (t/s)",
    "Uncached Throughput / box (t/s/hardware)",
    "Cached Throughput (t/s)",
    "Cached Throughput / box (t/s/hardware)",
    "TTFT (ms)",
    "Real Prompt Speed (t/s/user)",
    "Prompt Speed with Queueing (t/s/user)",
    "Gen Speed (t/s/user)",
    "RPM",
]

MalformedCase = Literal["missing-column", "invalid-number", "empty-table"]

# Model A and Model L share config keys but have different performance values.
_SCALE = {"Model A": 1.0, "Model L": 1.5}


def build_valid_workbook_bytes(model_name: str) -> bytes:
    """Return fresh .xlsx bytes for a deterministic two-row, profile-1 sweep."""
    if model_name not in _SCALE:
        raise ValueError(f"Unknown fixture model: {model_name}")
    s = _SCALE[model_name]

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.cell(row=1, column=1, value=f"{model_name} profile 1")  # title row, ignored by parser
    for col, header in enumerate(HEADERS, start=1):
        ws.cell(row=2, column=col, value=header)

    # Same keys for every model: input 1024, output 128, cache 0.5, batch 1 then 2.
    # Columns: in, out, cache, batch, max_ms, target_ms, prompt_tp, gen_tp,
    #          tp, tp/box, uncached, uncached/box, cached, cached/box,
    #          ttft, real_prompt, prompt_queued, gen_speed, rpm
    rows = [
        [
            1024,
            128,
            0.5,
            1,
            2000,
            1500,
            2000 * s,
            800 * s,
            1000 * s,
            250 * s,
            400 * s,
            100 * s,
            600 * s,
            150 * s,
            100 * s,
            900 * s,
            850 * s,
            40 * s,
            60 * s,
        ],
        [
            1024,
            128,
            0.5,
            2,
            2500,
            1500,
            3600 * s,
            1400 * s,
            1800 * s,
            450 * s,
            720 * s,
            180 * s,
            1080 * s,
            270 * s,
            140 * s,
            800 * s,
            760 * s,
            36 * s,
            108 * s,
        ],
    ]
    for r, values in enumerate(rows, start=3):
        for c, value in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=value)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_malformed_workbook_bytes(case: MalformedCase) -> bytes:
    """Mutate a fresh workbook to isolate one specific validation failure."""
    wb = load_workbook(io.BytesIO(build_valid_workbook_bytes("Model L")))
    ws = wb["Summary"]
    batch_col = HEADERS.index("Batch Size") + 1

    if case == "missing-column":
        ws.cell(row=2, column=batch_col).value = None
    elif case == "invalid-number":
        ws.cell(row=3, column=batch_col, value="not-a-number")
    elif case == "empty-table":
        ws.delete_rows(3, ws.max_row - 2)  # remove every data row, keep rows 1-2
    else:
        raise ValueError(f"Unknown malformed case: {case}")

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_browser_fixtures() -> None:
    """Export only test inputs; these are never frontend public/sample assets."""
    # Build everything before writing so unfinished factories leave no partial export.
    files = {
        "Model A profile 1.xlsx": build_valid_workbook_bytes("Model A"),
        "Model L profile 1.xlsx": build_valid_workbook_bytes("Model L"),
        "missing-column.xlsx": build_malformed_workbook_bytes("missing-column"),
        "invalid-number.xlsx": build_malformed_workbook_bytes("invalid-number"),
        "empty-table.xlsx": build_malformed_workbook_bytes("empty-table"),
        "corrupt.xlsx": b"This is not an Excel ZIP archive.",
    }
    destination = (
        Path(__file__).resolve().parents[2] / "frontend" / "tests" / "fixtures" / "generated"
    )
    destination.mkdir(parents=True, exist_ok=True)
    for filename, contents in files.items():
        (destination / filename).write_bytes(contents)
    print(f"Exported {len(files)} test fixtures to {destination}")


if __name__ == "__main__":
    export_browser_fixtures()

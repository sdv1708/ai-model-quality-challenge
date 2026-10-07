"""Issue #10 public-boundary tests."""

import io
from typing import Literal

import pytest
from fastapi.testclient import TestClient
from httpx import Response
from openpyxl import load_workbook  # type: ignore[import-untyped]
from resilience_fixtures import (
    MalformedCase,
    build_malformed_workbook_bytes,
    build_valid_workbook_bytes,
)

from perf_api.main import app

CLIENT = TestClient(app)
ENDPOINT = "/api/v1/comparisons/workbooks"
DECISION_ENDPOINT = "/api/v1/decisions/evaluate"
ENGINEERING_ENDPOINT = "/api/v1/engineering/analyze"

# Expected normalized values, taken from the fixture factory (scale: A = 1.0, L = 1.5).
EXPECTED = {
    "Model A": {
        1: {"throughput": 1000.0, "throughput_per_box": 250.0, "ttft_ms": 100.0, "gen_speed": 40.0},
        2: {"throughput": 1800.0, "throughput_per_box": 450.0, "ttft_ms": 140.0, "gen_speed": 36.0},
    },
    "Model L": {
        1: {"throughput": 1500.0, "throughput_per_box": 375.0, "ttft_ms": 150.0, "gen_speed": 60.0},
        2: {"throughput": 2700.0, "throughput_per_box": 675.0, "ttft_ms": 210.0, "gen_speed": 54.0},
    },
}


def _upload(*named: tuple[str, bytes]) -> Response:
    response: Response = CLIENT.post(ENDPOINT, files=[("workbooks", item) for item in named])
    return response


def _valid(model: str) -> tuple[str, bytes]:
    return (f"{model} profile 1.xlsx", build_valid_workbook_bytes(model))


def _with_second_batch(model: str, batch_size: int) -> tuple[str, bytes]:
    """A valid sweep whose second row uses another batch size (creates a coverage gap)."""
    wb = load_workbook(io.BytesIO(build_valid_workbook_bytes(model)))
    wb["Summary"].cell(row=4, column=4).value = batch_size
    buffer = io.BytesIO()
    wb.save(buffer)
    return (f"{model} profile 1.xlsx", buffer.getvalue())


def test_fresh_model_l_reaches_both_analysis_endpoints() -> None:
    response = _upload(_valid("Model L"))
    assert response.status_code == 200
    body = response.json()
    assert body["models"] == ["Model L"]
    assert body["workbooks"][0]["record_count"] == 2
    assert body["diagnostics"] == []

    workbook = body["workbooks"][0]
    assert workbook["model_name"] == "Model L"
    assert workbook["profile_id"] == "1"
    for record, batch in zip(workbook["records"], (1, 2), strict=True):
        expected = EXPECTED["Model L"][batch]
        assert record["input_length"] == 1024
        assert record["output_length"] == 128
        assert record["cache_percentage"] == 0.5
        assert record["batch_size"] == batch
        for field, value in expected.items():
            assert record[field] == value
        # Components stay coherent: cached + uncached == total.
        assert record["cached_throughput"] + record["uncached_throughput"] == record["throughput"]

    # Customer API: batch 1 fails the throughput target, batch 2 meets all targets.
    decision = CLIENT.post(
        DECISION_ENDPOINT,
        json={
            "workbook": workbook,
            "targets": {
                "input_tokens": 1024,
                "output_tokens": 128,
                "cache_fraction": 0.5,
                "min_throughput_tps": 2000,
                "min_generation_speed_tps_per_user": 50,
                "max_ttft_ms": 250,
            },
            "assumptions": {},
        },
    )
    assert decision.status_code == 200
    verdict = decision.json()
    assert verdict["status"] == "go"
    assert verdict["model_name"] == "Model L"
    assert verdict["profile_id"] == "1"
    assert verdict["selected_record"]["batch_size"] == 2
    actuals = {item["metric"]: item for item in verdict["evidence"]}
    assert actuals["throughput"]["actual"] == 2700.0
    assert actuals["throughput"]["target"] == 2000.0
    assert actuals["generation_speed"]["actual"] == 54.0
    assert actuals["ttft"]["actual"] == 210.0
    assert all(item["outcome"] == "met" for item in verdict["evidence"])
    batch_one = verdict["evaluated_configurations"][0]
    assert batch_one["record"]["batch_size"] == 1
    assert batch_one["unmet_constraints"] == ["throughput"]

    # Engineering API: source identity and exact numeric values survive.
    engineering = CLIENT.post(ENGINEERING_ENDPOINT, json={"workbooks": body["workbooks"]})
    assert engineering.status_code == 200
    analysis = engineering.json()
    assert [c["source"]["record_index"] for c in analysis["configurations"]] == [0, 1]
    for config, batch in zip(analysis["configurations"], (1, 2), strict=True):
        assert config["source"]["model_name"] == "Model L"
        assert config["source"]["profile_id"] == "1"
        assert config["source"]["workbook_index"] == 0
        metrics = {m["name"]: m for m in config["metrics"]}
        assert metrics["throughput"]["value"] == EXPECTED["Model L"][batch]["throughput"]
        assert metrics["throughput"]["unit"] == "t/s"
        assert (
            metrics["throughput_per_box"]["value"]
            == EXPECTED["Model L"][batch]["throughput_per_box"]
        )
        assert metrics["ttft_ms"]["unit"] == "ms"
    batch_trend = next(
        t
        for t in analysis["trends"]
        if t["dimension"] == "batch_size" and t["metric_name"] == "throughput"
    )
    assert batch_trend["baseline_value"] == 1500.0
    assert batch_trend["candidate_value"] == 2700.0
    assert batch_trend["baseline"]["record_index"] == 0
    assert batch_trend["candidate"]["record_index"] == 1
    assert analysis["anomalies"] == []


def test_fresh_models_align_independently_of_upload_order() -> None:
    forward = _upload(_valid("Model A"), _valid("Model L"))
    reverse = _upload(_valid("Model L"), _valid("Model A"))
    assert forward.status_code == reverse.status_code == 200
    assert forward.json() == reverse.json()

    body = forward.json()
    assert body["models"] == ["Model A", "Model L"]
    assert body["diagnostics"] == []
    assert [c["key"]["batch_size"] for c in body["configurations"]] == [1, 2]
    for config, batch in zip(body["configurations"], (1, 2), strict=True):
        assert config["is_comparable"] is True
        assert config["missing_models"] == []
        assert config["key"]["profile_id"] == "1"
        assert config["key"]["input_length"] == 1024
        assert config["key"]["output_length"] == 128
        assert config["key"]["cache_percentage"] == 0.5
        members = {m["model_name"]: m["record"] for m in config["members"]}
        assert list(members) == ["Model A", "Model L"]
        for model in ("Model A", "Model L"):
            for field, value in EXPECTED[model][batch].items():
                assert members[model][field] == value
        assert members["Model A"]["throughput"] != members["Model L"]["throughput"]


def test_missing_configuration_is_a_gap_not_a_zero() -> None:
    # A covers batches 1 and 3; L covers batches 1 and 2.
    response = _upload(_with_second_batch("Model A", 3), _valid("Model L"))
    assert response.status_code == 200
    configs = {c["key"]["batch_size"]: c for c in response.json()["configurations"]}
    assert sorted(configs) == [1, 2, 3]

    assert configs[1]["is_comparable"] is True
    assert configs[1]["missing_models"] == []

    assert configs[2]["is_comparable"] is False
    assert configs[2]["missing_models"] == ["Model A"]
    assert [m["model_name"] for m in configs[2]["members"]] == ["Model L"]

    assert configs[3]["is_comparable"] is False
    assert configs[3]["missing_models"] == ["Model L"]
    assert [m["model_name"] for m in configs[3]["members"]] == ["Model A"]
    assert configs[3]["members"][0]["record"]["throughput"] == 1800.0  # A's real value, not 0


@pytest.mark.parametrize(
    ("case", "expected_text"),
    [
        ("missing-column", "Batch Size"),
        ("invalid-number", "Row 3 numeric validation failed"),
        ("empty-table", "no performance data rows"),
        ("corrupt", "Invalid Excel workbook file"),
    ],
)
def test_unusable_batch_names_the_file_and_explains_the_failure(
    case: MalformedCase | Literal["corrupt"], expected_text: str
) -> None:
    contents = (
        b"This is not an Excel ZIP archive."
        if case == "corrupt"
        else build_malformed_workbook_bytes(case)
    )
    filename = f"{case}.xlsx"
    response = _upload((filename, contents))
    assert response.status_code == 422
    detail = response.json()["detail"]
    assert isinstance(detail, list)
    assert len(detail) == 1
    assert detail[0]["code"] == "PARSER_ERROR"
    assert detail[0]["filename"] == filename
    assert expected_text in detail[0]["message"]


def test_mixed_batch_retains_model_l_with_file_specific_diagnostics() -> None:
    response = CLIENT.post(
        ENDPOINT,
        files=[
            ("workbooks", ("Model L profile 1.xlsx", build_valid_workbook_bytes("Model L"))),
            (
                "workbooks",
                ("missing-column.xlsx", build_malformed_workbook_bytes("missing-column")),
            ),
        ],
    )
    assert response.status_code == 200
    body = response.json()
    assert body["models"] == ["Model L"]

    assert len(body["diagnostics"]) == 1
    diagnostic = body["diagnostics"][0]
    assert diagnostic["code"] == "PARSER_ERROR"
    assert diagnostic["filename"] == "missing-column.xlsx"
    assert "Batch Size" in diagnostic["message"]

    assert len(body["workbooks"]) == 1
    records = body["workbooks"][0]["records"]
    assert [r["batch_size"] for r in records] == [1, 2]
    for record, batch in zip(records, (1, 2), strict=True):
        for field, value in EXPECTED["Model L"][batch].items():
            assert record[field] == value
    assert [c["key"]["batch_size"] for c in body["configurations"]] == [1, 2]
    assert all(c["missing_models"] == [] for c in body["configurations"])

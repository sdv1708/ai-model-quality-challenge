"""Workbook parsing behavior at the public parser boundary."""

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook  # type: ignore[import-untyped]

from perf_api.parser import WorkbookParseError, parse_and_normalize_workbook


def test_populated_row_without_batch_size_is_rejected() -> None:
    """A missing required cell must not truncate the remaining sweep."""
    zip_path = Path(__file__).parents[2] / "perf_data.zip"
    with ZipFile(zip_path) as archive:
        workbook_bytes = archive.read("Model_A_profile_1/Model A profile 1.xlsx")

    workbook = load_workbook(BytesIO(workbook_bytes))
    workbook["Summary"]["D4"] = None
    modified = BytesIO()
    workbook.save(modified)

    with pytest.raises(WorkbookParseError, match=r"row 4.*Batch Size"):
        parse_and_normalize_workbook(modified.getvalue(), "Model A profile 1.xlsx")


def test_blank_row_does_not_hide_later_performance_rows() -> None:
    """A gap in the table must not turn the remainder into a successful upload."""
    zip_path = Path(__file__).parents[2] / "perf_data.zip"
    with ZipFile(zip_path) as archive:
        workbook_bytes = archive.read("Model_A_profile_1/Model A profile 1.xlsx")

    workbook = load_workbook(BytesIO(workbook_bytes))
    sheet = workbook["Summary"]
    for row in sheet.iter_rows(min_row=4, max_row=4, min_col=4, max_col=19):
        for cell in row:
            cell.value = None
    modified = BytesIO()
    workbook.save(modified)

    with pytest.raises(WorkbookParseError, match=r"row 4"):
        parse_and_normalize_workbook(modified.getvalue(), "Model A profile 1.xlsx")

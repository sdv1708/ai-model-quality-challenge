"""Tests for the workbook HTTP route."""

import zipfile
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook  # type: ignore[import-untyped]

from perf_api.main import app

CLIENT = TestClient(app)


@pytest.fixture
def sample_workbook_bytes() -> bytes:
    """Return one supplied workbook as an upload-ready byte string."""
    zip_path = Path(__file__).parents[2] / "perf_data.zip"
    with zipfile.ZipFile(zip_path) as archive:
        return archive.read("Model_A_profile_1/Model A profile 1.xlsx")


def test_workbook_normalization_rejects_non_xlsx() -> None:
    """Non-xlsx files return 400 Bad Request."""
    response = CLIENT.post(
        "/api/v1/workbooks/normalize",
        files={
            "workbook": (
                "example.txt",
                b"invalid file content",
                "text/plain",
            )
        },
    )
    assert response.status_code == 400
    assert "Only .xlsx workbook files are supported." in response.json()["detail"]


def test_workbook_normalization_parses_real_sweep(sample_workbook_bytes: bytes) -> None:
    """A valid xlsx sweep from perf_data.zip is normalized into typed records."""
    response = CLIENT.post(
        "/api/v1/workbooks/normalize",
        files={
            "workbook": (
                "Model A profile 1.xlsx",
                sample_workbook_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 200
    data = response.json()
    assert data["model_name"] == "Model A"
    assert data["profile_id"] == "1"
    assert data["record_count"] == 4
    # Check that forward-filling worked for Input Length, Output Length, Cache %
    first_record = data["records"][0]
    assert first_record["input_length"] == 10000
    assert first_record["output_length"] == 333
    assert first_record["cache_percentage"] == 0.5
    assert first_record["batch_size"] == 10

    fourth_record = data["records"][3]
    assert fourth_record["input_length"] == 10000
    assert fourth_record["output_length"] == 333
    assert fourth_record["cache_percentage"] == 0.5
    assert fourth_record["batch_size"] == 40


def test_workbook_normalization_accepts_unseen_model_name(sample_workbook_bytes: bytes) -> None:
    """A conforming upload discovers the model name from its filename."""
    response = CLIENT.post(
        "/api/v1/workbooks/normalize",
        files={"workbook": ("Model L profile 1.xlsx", sample_workbook_bytes)},
    )

    assert response.status_code == 200
    assert response.json()["model_name"] == "Model L"


def test_workbook_normalization_reports_missing_required_column(
    sample_workbook_bytes: bytes,
) -> None:
    """An absent source header produces a specific HTTP validation error."""
    workbook = load_workbook(BytesIO(sample_workbook_bytes))
    workbook["Summary"]["D2"] = None
    modified = BytesIO()
    workbook.save(modified)

    response = CLIENT.post(
        "/api/v1/workbooks/normalize",
        files={"workbook": ("Model A profile 1.xlsx", modified.getvalue())},
    )

    assert response.status_code == 422
    assert "Batch Size" in response.json()["detail"]


def test_workbook_normalization_reports_invalid_numeric_cell(
    sample_workbook_bytes: bytes,
) -> None:
    """A nonnumeric batch size identifies the source row and column."""
    workbook = load_workbook(BytesIO(sample_workbook_bytes))
    workbook["Summary"]["D3"] = "not-a-number"
    modified = BytesIO()
    workbook.save(modified)

    response = CLIENT.post(
        "/api/v1/workbooks/normalize",
        files={"workbook": ("Model A profile 1.xlsx", modified.getvalue())},
    )

    assert response.status_code == 422
    detail = response.json()["detail"]
    assert "Row 3" in detail
    assert "Batch Size" in detail

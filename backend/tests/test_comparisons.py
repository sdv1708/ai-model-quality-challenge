"""Tests for issue #6 multi-workbook comparison endpoint and alignment logic."""

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from fastapi.testclient import TestClient
from openpyxl import load_workbook  # type: ignore[import-untyped]

from perf_api.main import app

client = TestClient(app)


def get_workbook_bytes(name: str) -> bytes:
    zip_path = Path(__file__).parents[2] / "perf_data.zip"
    with ZipFile(zip_path) as archive:
        return archive.read(name)


def test_compare_single_workbook() -> None:
    wb_bytes = get_workbook_bytes("Model_A_profile_1/Model A profile 1.xlsx")
    response = client.post(
        "/api/v1/comparisons/workbooks",
        files=[
            (
                "workbooks",
                (
                    "Model A profile 1.xlsx",
                    wb_bytes,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            )
        ],
    )
    assert response.status_code == 200
    data = response.json()
    assert data["models"] == ["Model A"]
    assert len(data["workbooks"]) == 1
    assert len(data["configurations"]) > 0
    assert all(c["is_comparable"] is False for c in data["configurations"])


def test_compare_multiple_models_and_order_independence() -> None:
    wb_a = get_workbook_bytes("Model_A_profile_1/Model A profile 1.xlsx")
    wb_b = get_workbook_bytes("Model_B_profile_1/Model B profile 1.xlsx")

    res1 = client.post(
        "/api/v1/comparisons/workbooks",
        files=[
            ("workbooks", ("Model A profile 1.xlsx", wb_a, "application/octet-stream")),
            ("workbooks", ("Model B profile 1.xlsx", wb_b, "application/octet-stream")),
        ],
    )
    res2 = client.post(
        "/api/v1/comparisons/workbooks",
        files=[
            ("workbooks", ("Model B profile 1.xlsx", wb_b, "application/octet-stream")),
            ("workbooks", ("Model A profile 1.xlsx", wb_a, "application/octet-stream")),
        ],
    )

    assert res1.status_code == 200
    assert res2.status_code == 200
    data1 = res1.json()
    data2 = res2.json()

    assert data1 == data2
    assert data1["models"] == ["Model A", "Model B"]
    assert any(c["is_comparable"] is True for c in data1["configurations"])


def test_duplicate_sweep_diagnostic() -> None:
    wb_a = get_workbook_bytes("Model_A_profile_1/Model A profile 1.xlsx")
    res = client.post(
        "/api/v1/comparisons/workbooks",
        files=[
            ("workbooks", ("Model A profile 1.xlsx", wb_a, "application/octet-stream")),
            ("workbooks", ("Model A profile 1.xlsx", wb_a, "application/octet-stream")),
        ],
    )
    assert res.status_code == 200
    data = res.json()
    assert len(data["diagnostics"]) == 1
    assert data["diagnostics"][0]["code"] == "DUPLICATE_SWEEP"


def test_invalid_file_diagnostic() -> None:
    res = client.post(
        "/api/v1/comparisons/workbooks",
        files=[
            ("workbooks", ("invalid.txt", b"not an excel file", "text/plain")),
        ],
    )
    assert res.status_code == 422
    assert res.json()["detail"][0]["code"] == "INVALID_FILE_TYPE"
    assert res.json()["detail"][0]["filename"] == "invalid.txt"


def test_mixed_batch_preserves_valid_workbook_and_names_invalid_file() -> None:
    wb_a = get_workbook_bytes("Model_A_profile_1/Model A profile 1.xlsx")
    response = client.post(
        "/api/v1/comparisons/workbooks",
        files=[
            ("workbooks", ("bad.txt", b"bad", "text/plain")),
            ("workbooks", ("Model A profile 1.xlsx", wb_a, "application/octet-stream")),
        ],
    )
    assert response.status_code == 200
    body = response.json()
    assert body["models"] == ["Model A"]
    assert body["diagnostics"][0]["filename"] == "bad.txt"


def test_conflicting_duplicate_sweeps_do_not_depend_on_upload_order() -> None:
    original = get_workbook_bytes("Model_A_profile_1/Model A profile 1.xlsx")
    workbook = load_workbook(BytesIO(original))
    workbook["Summary"]["I3"] = 1
    changed = BytesIO()
    workbook.save(changed)
    wb_b = get_workbook_bytes("Model_B_profile_1/Model B profile 1.xlsx")

    def compare(first: bytes, second: bytes) -> dict[str, object]:
        response = client.post(
            "/api/v1/comparisons/workbooks",
            files=[
                ("workbooks", ("Model A profile 1.xlsx", first, "application/octet-stream")),
                ("workbooks", ("Model A profile 1.xlsx", second, "application/octet-stream")),
                ("workbooks", ("Model B profile 1.xlsx", wb_b, "application/octet-stream")),
            ],
        )
        assert response.status_code == 200
        return response.json()  # type: ignore[no-any-return]

    forward = compare(original, changed.getvalue())
    reverse = compare(changed.getvalue(), original)
    assert forward == reverse
    assert forward["models"] == ["Model B"]
    assert forward["diagnostics"][0]["code"] == "CONFLICTING_SWEEP"  # type: ignore[index]
    only_conflicts = client.post(
        "/api/v1/comparisons/workbooks",
        files=[
            ("workbooks", ("Model A profile 1.xlsx", original, "application/octet-stream")),
            (
                "workbooks",
                ("Model A profile 1.xlsx", changed.getvalue(), "application/octet-stream"),
            ),
        ],
    )
    assert only_conflicts.status_code == 422
    assert only_conflicts.json()["detail"][0]["code"] == "CONFLICTING_SWEEP"


def test_duplicate_configuration_excludes_ambiguous_sweep() -> None:
    original = get_workbook_bytes("Model_A_profile_1/Model A profile 1.xlsx")
    workbook = load_workbook(BytesIO(original))
    sheet = workbook["Summary"]
    for column in range(1, 20):
        sheet.cell(row=7, column=column).value = sheet.cell(row=3, column=column).value
    changed = BytesIO()
    workbook.save(changed)
    wb_b = get_workbook_bytes("Model_B_profile_1/Model B profile 1.xlsx")

    response = client.post(
        "/api/v1/comparisons/workbooks",
        files=[
            (
                "workbooks",
                ("Model A profile 1.xlsx", changed.getvalue(), "application/octet-stream"),
            ),
            ("workbooks", ("Model B profile 1.xlsx", wb_b, "application/octet-stream")),
        ],
    )
    assert response.status_code == 200
    body = response.json()
    assert any(item["code"] == "DUPLICATE_CONFIGURATION" for item in body["diagnostics"])
    assert body["models"] == ["Model B"]
    assert all(
        member["model_name"] == "Model B"
        for item in body["configurations"]
        for member in item["members"]
    )
    only_ambiguous = client.post(
        "/api/v1/comparisons/workbooks",
        files=[
            (
                "workbooks",
                ("Model A profile 1.xlsx", changed.getvalue(), "application/octet-stream"),
            )
        ],
    )
    assert only_ambiguous.status_code == 422
    assert only_ambiguous.json()["detail"][0]["code"] == "DUPLICATE_CONFIGURATION"


def test_unseen_model_and_missing_profile_are_explicit() -> None:
    wb_a = get_workbook_bytes("Model_A_profile_1/Model A profile 1.xlsx")
    wb_b = get_workbook_bytes("Model_B_profile_1/Model B profile 1.xlsx")
    wb_c = get_workbook_bytes("Model_C_profile_2/Model C profile 2.xlsx")
    response = client.post(
        "/api/v1/comparisons/workbooks",
        files=[
            ("workbooks", ("Model A profile 1.xlsx", wb_a, "application/octet-stream")),
            ("workbooks", ("Model B profile 1.xlsx", wb_b, "application/octet-stream")),
            ("workbooks", ("Model C profile 2.xlsx", wb_c, "application/octet-stream")),
            ("workbooks", ("Model L profile 1.xlsx", wb_a, "application/octet-stream")),
        ],
    )
    assert response.status_code == 200
    body = response.json()
    assert body["models"] == ["Model A", "Model B", "Model C", "Model L"]
    profile_two = [item for item in body["configurations"] if item["key"]["profile_id"] == "2"]
    assert profile_two
    assert all(item["is_comparable"] is False for item in profile_two)
    assert all("Model A" in item["missing_models"] for item in profile_two)


def test_incomplete_workbook_is_reported_by_filename() -> None:
    original = get_workbook_bytes("Model_A_profile_1/Model A profile 1.xlsx")
    workbook = load_workbook(BytesIO(original))
    workbook["Summary"]["D2"] = None
    changed = BytesIO()
    workbook.save(changed)
    response = client.post(
        "/api/v1/comparisons/workbooks",
        files=[
            (
                "workbooks",
                ("Model A profile 1.xlsx", changed.getvalue(), "application/octet-stream"),
            )
        ],
    )
    assert response.status_code == 422
    diagnostic = response.json()["detail"][0]
    assert diagnostic["filename"] == "Model A profile 1.xlsx"
    assert "Batch Size" in diagnostic["message"]
